import os
import re
import pymupdf
import requests
from django.conf import settings
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# ✅ Extract Text from PDF
def extract_text_from_pdf(file_path):
    try:
        doc = pymupdf.open(file_path)
        text = "\n".join([page.get_text("text") for page in doc])
        return text or "No text found."
    except Exception:
        return "Error extracting text."

# ✅ AI Summary Generation with Custom User Prompt
def summarize_with_deepseek(text, prompt="Summarize this document:"):
    print("🤖 Sending text to OpenRouter AI...")

    OPENROUTER_API_KEY = "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"
    headers = {"Authorization": f"Bearer {OPENROUTER_API_KEY}"}
    data = {
        "model": "deepseek/deepseek-r1-zero:free",
        "messages": [{"role": "user", "content": f"{prompt}\n{text}"} if text else {"role": "user", "content": f"{prompt}"}],
        "temperature": 0.7,
    }

    try:
        response = requests.post("https://openrouter.ai/api/v1/chat/completions", json=data, headers=headers)
        if response.status_code == 200:
            ai_response = response.json().get("choices", [{}])[0].get("message", {}).get("content", "Error generating summary.")
            print(f"✅ AI Response Received: {ai_response[:100]}...")
            return ai_response
        else:
            print(f"❌ OpenRouter Error: {response.text}")
            return "Error in AI processing."
    except Exception as e:
        print(f"❌ AI Request Failed: {e}")
        return "AI request failed."

def create_ppt(summary_text, ppt_path):
    try:
        prs = Presentation()
        slide_layout = prs.slide_layouts[1]  # Title and Content

        # Remove boxed markdown fences
        cleaned_text = re.sub(r"\\boxed{```markdown|```}", "", summary_text.strip(), flags=re.MULTILINE)

        # Split into individual slides
        slides = re.split(r'## Slide \d+: ', cleaned_text)
        cleaned_slides = []
        for slide in slides:
            slide = re.sub(r'^##?\s*Slide \d+ .*?\n?', '', slide.strip(), flags=re.MULTILINE)
            cleaned_slides.append(slide)

        for slide_data in cleaned_slides:
            if not slide_data.strip():
                continue

            # Extract title
            title_match = re.search(r'(?:\*\*|)Title(?:\*\*|):\s*(.*)', slide_data)
            slide_title = title_match.group(1).strip() if title_match else ""
            slide_title = re.sub(r"\*+", "", slide_title).strip()

            # Remove title from content
            content = slide_data
            if title_match:
                content = content.replace(title_match.group(0), "")

            # Extract table (markdown format)
            table_match = re.search(r'(\|.+\|(?:\n\|[-| ]+\|)+\n(?:\|.*\|\n?)+)', content)
            table_data = table_match.group(1).strip() if table_match else None
            if table_data:
                content = content.replace(table_data, "")

            # Process content lines for bullets
            lines = content.strip().splitlines()
            bullet_points = []
            parent = None

            for line in lines:
                stripped = line.strip()
                if stripped in ["#", "*", "-", "—"]:
                    continue

                stripped = re.sub(r'\$(.*?)\$', r'\1', stripped)
                if not stripped:
                    continue

                if stripped.endswith(":"):
                    parent = re.sub(r"\*+", "", stripped.strip("-• ")).strip()
                    bullet_points.append((parent, 0))
                elif parent:
                    bullet_points.append((stripped.strip("-• "), 1))
                else:
                    bullet_points.append((stripped.strip("-• "), 0))

            # Create slide
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.title.text = slide_title
            content_box = slide.placeholders[1]
            tf = content_box.text_frame
            tf.clear()

            has_bullets = bool(bullet_points)
            has_table = bool(table_data)

            # Decide vertical positioning
            bullet_top = Inches(1.5)
            table_top = Inches(1.5 if not has_bullets else 3.5)

            # Add bullets
            if has_bullets:
                for bullet, level in bullet_points:
                    if bullet:
                        p = tf.add_paragraph()
                        p.level = level
                        p.font.size = Pt(18)
                        p.clear()  # Clear default paragraph run

                        # Split by bold markers
                        parts = re.split(r'(\*\*.*?\*\*)', bullet)
                        for part in parts:
                            run = p.add_run()
                            clean_text = part.replace("**", "")
                            run.text = clean_text
                            run.font.size = Pt(18)
                            if part.startswith("**") and part.endswith("**"):
                                run.font.bold = True

            # Add table
            if has_table:
                rows = [row.strip() for row in table_data.strip().split("\n") if row.startswith("|")]
                table_rows = [row.strip("|").split("|") for row in rows]
                n_rows = len(table_rows)
                n_cols = len(table_rows[0])

                left = Inches(0.5)
                width = Inches(9)
                height = Inches(1.5)
                table_shape = slide.shapes.add_table(n_rows, n_cols, left, table_top, width, height).table

                for i, row in enumerate(table_rows):
                    for j, cell in enumerate(row):
                        table_shape.cell(i, j).text = cell.strip()
                        table_shape.cell(i, j).text_frame.paragraphs[0].font.size = Pt(12)

        prs.save(ppt_path)
        print(f"✅ PPT saved to: {ppt_path}")

    except Exception as e:
        print(f"❌ PPT Creation Error: {e}")